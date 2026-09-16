import json
from django.shortcuts import render, redirect, get_object_or_404
from django.contrib.auth.decorators import login_required
from django.contrib import messages
from django.db.models import Sum
from django.http import HttpResponse
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from .models import Income


@login_required
def income_view(request):
    if request.method == 'POST':
        source = request.POST.get('source')
        amount = request.POST.get('amount')
        date = request.POST.get('date')
        notes = request.POST.get('notes', '')

        if source and amount and date:
            Income.objects.create(
                user=request.user,
                source=source,
                amount=amount,
                date=date,
                notes=notes
            )
            messages.success(request, 'Income added successfully!')
            return redirect('income')
        else:
            messages.error(request, 'Please fill all required fields.')

    incomes = Income.objects.filter(user=request.user)
    total = sum(i.amount for i in incomes)

    # ---- Chart data: total amount grouped by source ----
    by_source = (
        incomes.values('source')
        .annotate(total=Sum('amount'))
        .order_by('-total')
    )
    source_labels = [dict(Income.SOURCE_CHOICES).get(row['source'], row['source']) for row in by_source]
    source_totals = [float(row['total']) for row in by_source]

    # ---- Chart data: total amount grouped by month (trend) ----
    by_month = (
        incomes.values('date__year', 'date__month')
        .annotate(total=Sum('amount'))
        .order_by('date__year', 'date__month')
    )
    month_labels = [f"{row['date__month']:02d}/{row['date__year']}" for row in by_month]
    month_totals = [float(row['total']) for row in by_month]

    context = {
        'incomes': incomes,
        'total': total,
        'source_choices': Income.SOURCE_CHOICES,
        'source_labels': json.dumps(source_labels),
        'source_totals': json.dumps(source_totals),
        'month_labels': json.dumps(month_labels),
        'month_totals': json.dumps(month_totals),
    }
    return render(request, 'income/income.html', context)


@login_required
def delete_income(request, pk):
    income = get_object_or_404(Income, pk=pk, user=request.user)
    income.delete()
    messages.success(request, 'Income deleted.')
    return redirect('income')


@login_required
def edit_income(request, pk):
    income = get_object_or_404(Income, pk=pk, user=request.user)

    if request.method == 'POST':
        source = request.POST.get('source')
        amount = request.POST.get('amount')
        date = request.POST.get('date')
        notes = request.POST.get('notes', '')

        if source and amount and date:
            income.source = source
            income.amount = amount
            income.date = date
            income.notes = notes
            income.save()
            messages.success(request, 'Income updated successfully!')
            return redirect('income')
        else:
            messages.error(request, 'Please fill all required fields.')

    context = {
        'income': income,
        'source_choices': Income.SOURCE_CHOICES,
    }
    return render(request, 'income/edit_income.html', context)


@login_required
def export_income_excel(request):
    incomes = Income.objects.filter(user=request.user)

    wb = Workbook()
    ws = wb.active
    ws.title = "Income"

    headers = ['Source', 'Amount', 'Date', 'Notes']
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill(start_color="4CAF50", end_color="4CAF50", fill_type="solid")

    total = 0
    for item in incomes:
        ws.append([item.get_source_display(), float(item.amount), item.date.strftime('%d-%m-%Y'), item.notes])
        total += item.amount

    ws.append([])
    ws.append(['Total', float(total)])
    ws['A' + str(ws.max_row)].font = Font(bold=True)
    ws['B' + str(ws.max_row)].font = Font(bold=True)

    for col in ws.columns:
        max_len = max(len(str(c.value)) if c.value else 0 for c in col)
        ws.column_dimensions[col[0].column_letter].width = max_len + 4

    response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    response['Content-Disposition'] = 'attachment; filename="income_report.xlsx"'
    wb.save(response)
    return response
