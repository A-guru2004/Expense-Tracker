import json
from django.shortcuts import render, redirect, get_object_or_404
from django.contrib.auth.decorators import login_required
from django.contrib import messages
from django.db.models import Sum
from django.http import HttpResponse
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from .models import Expense


@login_required
def expenses_view(request):
    if request.method == 'POST':
        category = request.POST.get('category')
        amount = request.POST.get('amount')
        date = request.POST.get('date')
        notes = request.POST.get('notes', '')

        if category and amount and date:
            Expense.objects.create(
                user=request.user,
                category=category,
                amount=amount,
                date=date,
                notes=notes
            )
            messages.success(request, 'Expense added successfully!')
            return redirect('expenses')
        else:
            messages.error(request, 'Please fill all required fields.')

    expenses = Expense.objects.filter(user=request.user)
    total = sum(e.amount for e in expenses)

    # ---- Chart data: total amount grouped by category ----
    by_category = (
        expenses.values('category')
        .annotate(total=Sum('amount'))
        .order_by('-total')
    )
    category_labels = [dict(Expense.CATEGORY_CHOICES).get(row['category'], row['category']) for row in by_category]
    category_totals = [float(row['total']) for row in by_category]

    # ---- Chart data: total amount grouped by month (trend) ----
    by_month = (
        expenses.values('date__year', 'date__month')
        .annotate(total=Sum('amount'))
        .order_by('date__year', 'date__month')
    )
    month_labels = [f"{row['date__month']:02d}/{row['date__year']}" for row in by_month]
    month_totals = [float(row['total']) for row in by_month]

    context = {
        'expenses': expenses,
        'total': total,
        'category_choices': Expense.CATEGORY_CHOICES,
        'category_labels': json.dumps(category_labels),
        'category_totals': json.dumps(category_totals),
        'month_labels': json.dumps(month_labels),
        'month_totals': json.dumps(month_totals),
    }
    return render(request, 'expenses/expenses.html', context)


@login_required
def delete_expense(request, pk):
    expense = get_object_or_404(Expense, pk=pk, user=request.user)
    expense.delete()
    messages.success(request, 'Expense deleted.')
    return redirect('expenses')


@login_required
def edit_expense(request, pk):
    expense = get_object_or_404(Expense, pk=pk, user=request.user)

    if request.method == 'POST':
        category = request.POST.get('category')
        amount = request.POST.get('amount')
        date = request.POST.get('date')
        notes = request.POST.get('notes', '')

        if category and amount and date:
            expense.category = category
            expense.amount = amount
            expense.date = date
            expense.notes = notes
            expense.save()
            messages.success(request, 'Expense updated successfully!')
            return redirect('expenses')
        else:
            messages.error(request, 'Please fill all required fields.')

    context = {
        'expense': expense,
        'category_choices': Expense.CATEGORY_CHOICES,
    }
    return render(request, 'expenses/edit_expense.html', context)


@login_required
def export_expense_excel(request):
    expenses = Expense.objects.filter(user=request.user)

    wb = Workbook()
    ws = wb.active
    ws.title = "Expenses"

    headers = ['Category', 'Amount', 'Date', 'Notes']
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill(start_color="EB5757", end_color="EB5757", fill_type="solid")

    total = 0
    for item in expenses:
        ws.append([item.get_category_display(), float(item.amount), item.date.strftime('%d-%m-%Y'), item.notes])
        total += item.amount

    ws.append([])
    ws.append(['Total', float(total)])
    ws['A' + str(ws.max_row)].font = Font(bold=True)
    ws['B' + str(ws.max_row)].font = Font(bold=True)

    for col in ws.columns:
        max_len = max(len(str(c.value)) if c.value else 0 for c in col)
        ws.column_dimensions[col[0].column_letter].width = max_len + 4

    response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    response['Content-Disposition'] = 'attachment; filename="expense_report.xlsx"'
    wb.save(response)
    return response
